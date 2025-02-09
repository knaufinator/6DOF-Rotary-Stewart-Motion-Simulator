import openpyxl
import numpy as np

def analyze_kinematics():
    # Load the workbook
    wb = openpyxl.load_workbook('documentation/StewartPlatformSimulator.xlsm', read_only=False, data_only=True)
    
    # Print available sheets
    print("Available sheets:", wb.sheetnames)
    
    # Analyze the DATA sheet
    data_sheet = wb['DATA']
    print("\nAnalyzing cells in DATA sheet:")
    
    # Print all non-empty cells in first 20 rows to see what we have
    for row in range(1, 21):
        row_data = []
        for col in range(1, 11):
            cell = data_sheet.cell(row=row, column=col)
            if cell.value is not None:
                row_data.append(f"{cell.coordinate}:{cell.value}")
        if row_data:
            print(f"Row {row}: {row_data}")
    
    # Also check the formulas (not just values)
    wb_formulas = openpyxl.load_workbook('documentation/StewartPlatformSimulator.xlsm', read_only=False, data_only=False)
    data_sheet_formulas = wb_formulas['DATA']
    
    print("\nAnalyzing formulas in DATA sheet:")
    for row in range(1, 21):
        for col in range(1, 11):
            cell = data_sheet_formulas.cell(row=row, column=col)
            if cell.value is not None and isinstance(cell.value, str) and cell.value.startswith('='):
                print(f"{cell.coordinate}: {cell.value}")

if __name__ == "__main__":
    analyze_kinematics()
